"""Tests for the scanned-PDF transcription wiring (no API spend).

Covers: invisible text layer synthesis on scanned pages (PDF side),
visible transcription section (DOCX side), scan image tagged as Artifact,
and the pdf2docx "Figure" placeholder alt-text fix.
"""

import os
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))
os.chdir(Path(__file__).parent.parent)

from ada_remediate import (
    convert_to_docx,
    make_docx_accessible,
    page_is_scanned,
    tag_pdf_with_accessibility,
)

ROOT = Path(__file__).parent.parent


def make_scanned_pdf(path: Path, pages: int = 2) -> None:
    """Create an image-only PDF (no text layer) — a synthetic scanned document."""
    import fitz

    doc = fitz.open()
    for _ in range(pages):
        page = doc.new_page(width=612, height=792)
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 120, 120))
        pix.set_rect(pix.irect, (180, 120, 60))
        page.insert_image(fitz.Rect(50, 50, 400, 300), pixmap=pix)
    doc.save(str(path))
    doc.close()


def fake_pages_data() -> list[dict]:
    """Vision-like elements for two scanned pages (no bboxes, like real vision)."""
    return [
        {
            "page": 1,
            "elements": [
                {"type": "heading", "level": 1, "text": "Additional Taxes on Qualified Plans"},
                {"type": "paragraph", "text": "Use this form to report additional taxes owed."},
                {
                    "type": "table",
                    "rows": [
                        {"is_header": True, "cells": ["Description", "Amount"]},
                        {"is_header": False, "cells": ["Early distribution", "1000"]},
                    ],
                },
                {"type": "list", "items": ["First requirement", "Second requirement"]},
            ],
        },
        {
            "page": 2,
            "elements": [
                {"type": "paragraph", "text": "Second page transcribed content here."},
            ],
        },
    ]


@pytest.fixture()
def scanned_pdf(tmp_path: Path) -> Path:
    pdf = tmp_path / "scanned_test.pdf"
    make_scanned_pdf(pdf)
    assert page_is_scanned(str(pdf), 1)
    assert page_is_scanned(str(pdf), 2)
    return pdf


def _struct_tag_names(pdf_path: Path) -> list[str]:
    """Collect /S tag names from the struct tree in document order."""
    import pikepdf

    tags: list[str] = []

    def walk(elem):
        try:
            s = str(elem.get("/S", ""))
        except Exception:
            return
        if s:
            tags.append(s)
        try:
            kids = elem.get("/K", [])
        except Exception:
            return
        if isinstance(kids, pikepdf.Array):
            for k in kids:
                try:
                    walk(k)
                except Exception:
                    continue

    with pikepdf.open(str(pdf_path)) as pdf:
        root = pdf.Root.get("/StructTreeRoot")
        assert root is not None, "no StructTreeRoot"
        kids = root.get("/K", [])
        for k in kids:
            walk(k)
    return tags


def test_scanned_pdf_gets_invisible_text_layer(scanned_pdf: Path, tmp_path: Path):
    out = tmp_path / "fixed.pdf"
    tag_pdf_with_accessibility(
        str(scanned_pdf), fake_pages_data(), str(out), scanned_pages={1, 2}
    )

    import fitz

    doc = fitz.open(str(out))
    assert len(doc) == 2
    p1 = doc[0].get_text().strip()
    p2 = doc[1].get_text().strip()
    # Transcription is extractable now (was 0 chars before the wiring fix)
    assert "Additional Taxes on Qualified Plans" in p1
    assert "Early distribution" in p1
    assert "First requirement" in p1
    assert "Second page transcribed content here." in p2
    # Reading order preserved: heading before paragraph before table
    assert p1.index("Additional Taxes") < p1.index("additional taxes owed")
    assert p1.index("additional taxes owed") < p1.index("Early distribution")
    doc.close()


def test_scanned_pdf_struct_tree_and_artifact_image(scanned_pdf: Path, tmp_path: Path):
    out = tmp_path / "fixed.pdf"
    tag_pdf_with_accessibility(
        str(scanned_pdf), fake_pages_data(), str(out), scanned_pages={1, 2}
    )
    tags = _struct_tag_names(out)
    # Content is tagged in reading order: heading, paragraph, table, list
    assert "/H1" in tags
    assert "/P" in tags
    assert "/Table" in tags
    assert "/L" in tags
    # The scan image itself must NOT be tagged as a Figure — it is the visual
    # layer and must be an Artifact (standard OCR-PDF practice).
    assert "/Figure" not in tags


def test_scanned_pdf_self_detect_fallback(scanned_pdf: Path, tmp_path: Path):
    """Without an explicit scanned_pages set, pages with no text blocks but
    vision elements are still detected as scanned."""
    out = tmp_path / "fixed.pdf"
    tag_pdf_with_accessibility(str(scanned_pdf), fake_pages_data(), str(out))
    import fitz

    doc = fitz.open(str(out))
    assert "Additional Taxes on Qualified Plans" in doc[0].get_text()
    doc.close()


def test_scanned_docx_gets_transcription_section(scanned_pdf: Path, tmp_path: Path):
    from docx import Document

    out = tmp_path / "fixed.docx"
    convert_to_docx(
        str(scanned_pdf),
        str(out),
        fake_pages_data(),
        page_dims=[(8.5, 11.0), (8.5, 11.0)],
        scanned_pages={1, 2},
    )
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    # Transcription sections appended (visible, per Jason's choice)
    assert any("Scanned page 1" in t and "transcribed text" in t for t in texts)
    assert any("Scanned page 2" in t and "transcribed text" in t for t in texts)
    # Content present as real structured content
    assert any("Additional Taxes on Qualified Plans" in t for t in texts)
    assert any("Second page transcribed content here." in t for t in texts)
    assert any("First requirement" in t for t in texts)
    # Table rendered as a real Word table
    assert len(doc.tables) >= 1
    assert any("Early distribution" in c.text for t in doc.tables for r in t.rows for c in r.cells)
    # Section heading uses a Heading style
    headings = [p for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any("Scanned page 1" in p.text for p in headings)


def _docx_with_placeholder_figure(tmp_path: Path) -> Path:
    """Minimal DOCX with one image whose descr is pdf2docx's 'Figure' placeholder."""
    from docx import Document
    from docx.shared import Inches
    import io

    # 1x1 red PNG
    import struct
    import zlib

    def png_chunk(typ: bytes, data: bytes) -> bytes:
        c = typ + data
        return struct.pack(">I", len(data)) + c + struct.pack(">I", zlib.crc32(c))

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    raw = b"\x00\xff\x00\x00"
    png = (
        b"\x89PNG\r\n\x1a\n"
        + png_chunk(b"IHDR", ihdr)
        + png_chunk(b"IDAT", zlib.compress(raw))
        + png_chunk(b"IEND", b"")
    )

    doc = Document()
    doc.add_paragraph("before")
    run = doc.add_paragraph().add_run()
    run.add_picture(io.BytesIO(png), width=Inches(1))
    # Simulate pdf2docx: pre-fill the placeholder descr
    drawing = run._r.find(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline"
    )
    docPr = drawing.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"
    )
    docPr.set("descr", "Figure")
    p = tmp_path / "fig.docx"
    doc.save(str(p))
    return p


def test_figure_placeholder_gets_vision_alt_text(tmp_path: Path):
    from docx import Document

    p = _docx_with_placeholder_figure(tmp_path)
    pages_data = [
        {"page": 1, "elements": [{"type": "image", "alt_text": "Bar chart of quarterly revenue"}]}
    ]
    make_docx_accessible(str(p), pages_data)
    doc = Document(str(p))
    ns = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    docPr = doc.element.body.find(f".//{{{ns}}}docPr")
    assert docPr.get("descr") == "Bar chart of quarterly revenue"


def test_figure_placeholder_without_alt_marked_decorative(tmp_path: Path):
    from docx import Document

    p = _docx_with_placeholder_figure(tmp_path)
    make_docx_accessible(str(p), [{"page": 1, "elements": []}])
    doc = Document(str(p))
    ns = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    docPr = doc.element.body.find(f".//{{{ns}}}docPr")
    assert docPr.get("decorative") == "1"
    assert (docPr.get("descr") or "") != "Figure"
